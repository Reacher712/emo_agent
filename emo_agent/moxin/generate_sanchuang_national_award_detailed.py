from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def set_layout(doc: Document):
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(2.5)
        sec.header_distance = Cm(1.5)
        sec.footer_distance = Cm(1.5)

        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.text = "三创赛项目计划书（国奖冲刺·详细版）"
        if hp.runs:
            run = hp.runs[0]
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)

        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.text = "第 "
        add_page_number(fp)
        fp.add_run(" 页")
        for run in fp.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)


def set_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(0.74)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "黑体"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h1.font.size = Pt(16)
    h1.font.bold = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "黑体"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h2.font.size = Pt(14)
    h2.font.bold = True


def add_para(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph(text)
    if p.runs:
        r = p.runs[0]
        r.bold = True
        r.font.name = "黑体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if p.runs:
        run = p.runs[0]
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("三创赛项目计划书（国奖冲刺·详细版）")
    r.bold = True
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    r.font.size = Pt(24)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("墨韵心调——AI驱动的国画多模态情感疗愈平台")
    r2.font.name = "宋体"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r2.font.size = Pt(15)

    doc.add_paragraph("")
    for line in [
        "版本定位：创新-验证-可复制导向",
        "适用场景：三创赛国奖评审与终审答辩",
        "日期：2026年3月",
    ]:
        pp = doc.add_paragraph(line)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_toc(doc: Document):
    doc.add_heading("目录", level=1)
    toc = [
        "第一章：执行摘要",
        "第二章：项目背景与问题定义",
        "第三章：市场需求与用户证据",
        "第四章：竞争格局与机会窗口",
        "第五章：产品与服务方案",
        "第六章：技术架构与实现路径",
        "第七章：创新点与知识产权布局（前置）",
        "第八章：实证验证与阶段成果（重点）",
        "第九章：运营模式与推广策略",
        "第十章：组织架构与实施计划",
        "第十一章：风险、伦理与合规治理（前置）",
        "第十二章：财务预测与分析",
        "第十三章：综合回报分析（经济+社会+教育）",
        "第十四章：发展规划与战略部署",
        "第十五章：结论与合作诉求",
        "附录A：市场调研详细数据",
        "附录B：技术方案详细文档",
    ]
    for line in toc:
        doc.add_paragraph(line)
    doc.add_page_break()


def main():
    doc = Document()
    set_layout(doc)
    set_styles(doc)
    add_cover(doc)
    add_toc(doc)

    doc.add_heading("第一章：执行摘要", level=1)
    add_subtitle(doc, "（一）项目定位")
    add_para(doc, "“墨韵心调”是面向青年群体情绪压力管理与关系沟通需求的多模态情感疗愈项目。项目以国画为文化主线，以人工智能为技术引擎，以数据追踪为评估依据，构建“识别—干预—反馈—优化”闭环服务。")
    add_subtitle(doc, "（二）核心问题")
    add_para(doc, "项目聚焦三类现实问题：第一，现有专业心理服务触达频率与日常陪伴能力不足；第二，通用情绪工具在文化适配与个体反馈方面存在缺口；第三，AI内容工具缺乏效果验证与长期追踪机制。")
    add_subtitle(doc, "（三）阶段成果")
    add_para(doc, "项目已形成总览方案、问卷数据脚本、可视化看板、竞品证据页与路演材料，完成需求调研样本（n=180）和试用评估样本（n=60）的阶段性积累。初步结果显示项目路径具有可行性与推广潜力。")

    doc.add_heading("第二章：项目背景与问题定义", level=1)
    add_subtitle(doc, "（一）政策与行业背景")
    add_para(doc, "心理健康教育体系建设、美育融合与数字化治理持续推进，为本项目提供了政策环境。与此同时，生成式AI技术发展提升了个性化内容生成能力，推动“心理支持+文化内容+数据反馈”融合成为可能。")
    add_subtitle(doc, "（二）问题链条界定")
    add_para(doc, "从用户路径看，青年群体在“产生压力—识别压力—寻求帮助—持续改善”过程中普遍存在断点。传统服务偏重后端干预，缺少前端日常支持；通用工具偏重即时体验，缺少长期效果评价。")
    add_subtitle(doc, "（三）项目意义")
    add_para(doc, "项目在理论层面探索国画情感量化机制，在技术层面探索多智能体协同路径，在应用层面构建校园可复制示范场景，具备学术价值、社会价值与教育价值。")
    add_caption(doc, "图2-1 项目问题链与价值链映射图（占位，待插入）")

    doc.add_heading("第三章：市场需求与用户证据", level=1)
    add_subtitle(doc, "（一）调研设计")
    add_para(doc, "项目采用“问卷+访谈+试用反馈”的混合调研设计。问卷侧重需求强度与功能偏好，访谈侧重行为动机与阻碍因素，试用反馈侧重可用性与感知效果。")
    add_subtitle(doc, "（二）关键数据证据")
    add_para(doc, "现有数据中，近1个月压力较高及以上占比67.2%，显示高频轻干预需求具有普遍性；愿意尝试匿名线上支持占比77.2%，说明用户对隐私友好型服务有较高接受度；愿意尝试艺术化干预占比81.7%，验证项目文化路径的适配性。")
    add_subtitle(doc, "（三）需求分层与优先级")
    add_para(doc, "从需求分层看，第一优先为个性化情绪干预（国画+音乐）；第二优先为趋势看板与周报反馈；第三优先为关系场景沟通支持。该优先级为产品迭代顺序与资源配置提供依据。")
    add_caption(doc, "表3-1 调研样本结构及核心指标（占位，待插入）")
    add_caption(doc, "图3-1 需求优先级雷达图（占位，待插入）")

    doc.add_heading("第四章：竞争格局与机会窗口", level=1)
    add_subtitle(doc, "（一）竞品类型划分")
    add_para(doc, "目前可比产品主要包括冥想类APP、心理咨询平台与AI生成工具。三类产品分别在轻量陪伴、专业支持和内容生成方面具备优势，但普遍缺乏“文化疗愈+追踪评估+关系场景”的系统闭环。")
    add_subtitle(doc, "（二）差异化价值")
    add_para(doc, "本项目的核心差异化在于：以国画为高识别度文化媒介，以多智能体协同保障内容质量，以数据追踪验证干预效果，以双人共绘拓展关系沟通场景。")
    add_subtitle(doc, "（三）机会窗口判断")
    add_para(doc, "在心理健康需求上升、校园场景数字化建设增强、文化内容创新需求增加的叠加背景下，项目具备形成“先场景验证、再规模复制”的窗口期。")
    add_caption(doc, "图4-1 竞品定位矩阵图（占位，待插入）")

    doc.add_heading("第五章：产品与服务方案", level=1)
    add_subtitle(doc, "（一）产品定位")
    add_para(doc, "产品定位为“面向青年日常情绪支持的文化型数字疗愈系统”，强调辅助性、连续性与可解释性，不替代医疗诊疗。")
    add_subtitle(doc, "（二）服务结构")
    add_para(doc, "服务结构分为个人疗愈与关系疗愈两大模块。个人疗愈模块包含情绪识别、国画生成、音乐匹配、看板与周报；关系疗愈模块包含双人共绘、行为分析与关系温度报告。")
    add_subtitle(doc, "（三）用户价值闭环")
    add_para(doc, "用户在使用过程中可获得“即时调节效果+趋势反馈结果+下一步行动建议”，从而实现从单次体验到长期使用的行为转化。")
    add_caption(doc, "图5-1 产品服务全景图（占位，待插入）")
    add_caption(doc, "表5-1 模块输入-输出-指标对照表（占位，待插入）")

    doc.add_heading("第六章：技术架构与实现路径", level=1)
    add_subtitle(doc, "（一）系统架构")
    add_para(doc, "技术架构由用户层、协调层、智能体层与数据层构成。协调层负责任务编排与状态管理，智能体层负责情感感知、艺术生成与音乐匹配，数据层负责过程存储与趋势分析。")
    add_subtitle(doc, "（二）关键机制")
    add_para(doc, "关键机制包括：多智能体协同机制、情感验证回路机制、动态追踪反馈机制。三者共同保证内容生成与目标情绪的一致性，并提升系统迭代效率。")
    add_subtitle(doc, "（三）实现路径")
    add_para(doc, "实现路径按“能力验证—模块联调—场景试点—数据复盘”推进。每一阶段均设置功能与指标双重验收，确保技术进展与业务目标一致。")
    add_caption(doc, "图6-1 系统技术架构图（占位，待插入）")
    add_caption(doc, "图6-2 核心处理流程图（占位，待插入）")

    doc.add_heading("第七章：创新点与知识产权布局（前置）", level=1)
    add_subtitle(doc, "（一）创新点体系")
    add_para(doc, "创新点一为协同式多智能体架构，突破单模型串行调用；创新点二为国画情感量化映射，提升内容可解释性；创新点三为动态情绪轨迹追踪，形成纵向评估能力；创新点四为双人共绘关系分析，拓展疗愈外延。")
    add_subtitle(doc, "（二）技术壁垒分析")
    add_para(doc, "项目壁垒不在单一模型参数，而在“多模块协同机制+情感验证方法+场景数据积累+反馈迭代系统”形成的复合能力。该复合能力具备较强迁移与复制价值。")
    add_subtitle(doc, "（三）知识产权策略")
    add_para(doc, "建议以软件著作权为基础，以流程方法文档和特色模块实现为补充，形成“成果归档—版本留痕—权属明确”的知识产权管理闭环。")
    add_caption(doc, "表7-1 创新点-壁垒-知识产权映射表（占位，待插入）")

    doc.add_heading("第八章：实证验证与阶段成果（重点）", level=1)
    add_subtitle(doc, "（一）验证目标与方法")
    add_para(doc, "本章目标是验证项目在可用性、有效性和可持续性三维度的可行性。方法上采用前后测设计与体验问卷结合，指标体系包括DASS-21、主观平静感、情绪可控感、SUS与UEQ等。")
    add_subtitle(doc, "（二）样本与口径说明")
    add_para(doc, "需求调研样本为180，试用评估样本为60。为保证结论可信度，所有结果需标注样本范围、采集时间和统计口径，避免不同章节口径不一致造成评审质疑。")
    add_subtitle(doc, "（三）核心结果解释")
    add_para(doc, "阶段结果显示，压力、焦虑、抑郁等指标呈下降趋势，平静感与可控感提升，SUS处于良好区间。该结果表明方案在“用户接受度—使用可行性—短期效果”方面具备正向信号。")
    add_subtitle(doc, "（四）阶段成果清单")
    add_para(doc, "项目已形成数据脚本、可视化看板、竞品证据页、路演PPT和计划书等多类成果，具备评审展示与后续扩展基础。")
    add_subtitle(doc, "（五）局限与改进计划")
    add_para(doc, "当前局限主要在样本规模、长期追踪时长和群体覆盖广度。下一阶段将通过扩大样本、延长追踪周期和引入外部评估提升结论稳健性。")
    add_caption(doc, "表8-1 验证指标体系与口径说明（占位，待插入）")
    add_caption(doc, "图8-1 DASS前后测结果图（占位，待插入）")
    add_caption(doc, "图8-2 可用性与体验结果图（占位，待插入）")
    add_caption(doc, "表8-2 阶段成果与证据链清单（占位，待插入）")

    doc.add_heading("第九章：运营模式与推广策略", level=1)
    add_subtitle(doc, "（一）运营模式")
    add_para(doc, "项目采用“校内试点—跨校复制—机构合作”分层运营模式。前期强调示范性与可验证性，中后期强调标准化与可复制性。")
    add_subtitle(doc, "（二）推广策略")
    add_para(doc, "推广采取“场景体验+内容传播+专家背书”策略，围绕核心证据构建传播链，避免泛化宣传。")
    add_subtitle(doc, "（三）关键运营指标")
    add_para(doc, "建议核心指标包含7日复访率、单次停留时长、报告查看率、共绘完成率和满意度，并形成月度复盘机制。")
    add_caption(doc, "图9-1 运营与推广路径图（占位，待插入）")

    doc.add_heading("第十章：组织架构与实施计划", level=1)
    add_subtitle(doc, "（一）团队结构")
    add_para(doc, "团队由技术、产品、心理、艺术与运营方向组成，强调跨学科协作。各模块设定负责人、验收标准和协同接口，保证执行效率。")
    add_subtitle(doc, "（二）实施计划")
    add_para(doc, "实施按阶段推进：准备期（数据与方案）、联调期（系统集成）、试点期（场景验证）、优化期（复盘迭代）。每阶段均配置“任务清单+交付物+量化指标”。")
    add_subtitle(doc, "（三）项目治理机制")
    add_para(doc, "建立周会与月度评估机制，形成“问题发现—责任分配—处理闭环—经验沉淀”的治理链条，提升项目稳定推进能力。")
    add_caption(doc, "图10-1 团队组织结构图（占位，待插入）")
    add_caption(doc, "表10-1 阶段里程碑与交付项（占位，待插入）")

    doc.add_heading("第十一章：风险、伦理与合规治理（前置）", level=1)
    add_subtitle(doc, "（一）风险分类")
    add_para(doc, "项目风险可分为技术风险、数据风险、运营风险和合规风险。其中特别关注心理相关场景中的误判、过度依赖、边界误解等风险。")
    add_subtitle(doc, "（二）合规机制")
    add_para(doc, "合规机制包括最小化采集、匿名化存储、分级权限、日志审计、内容风控和人工复核。项目明确非医疗诊断属性，避免功能边界外延。")
    add_subtitle(doc, "（三）高风险场景处置")
    add_para(doc, "针对高风险情绪表达场景，系统应触发风险提示与人工转介建议，并记录处置链路，确保使用安全与责任闭环。")
    add_caption(doc, "表11-1 风险与合规治理矩阵（占位，待插入）")

    doc.add_heading("第十二章：财务预测与分析", level=1)
    add_subtitle(doc, "（一）财务模型设定")
    add_para(doc, "财务预测采用保守、中性、积极三情景模型，围绕研发投入、算力运维、运营推广和服务收入四大维度进行测算。")
    add_subtitle(doc, "（二）成本结构分析")
    add_para(doc, "成本结构以研发与试点验证为主，前期投入集中于功能稳定与样本扩展，中后期投入逐步转向推广复制和合作服务。")
    add_subtitle(doc, "（三）稳健性分析")
    add_para(doc, "建议通过敏感性分析检验关键参数变动对结果的影响，以提高财务判断可信度和评审说服力。")
    add_caption(doc, "表12-1 财务预测参数与情景设定（占位，待插入）")

    doc.add_heading("第十三章：综合回报分析（经济+社会+教育）", level=1)
    add_subtitle(doc, "（一）经济回报")
    add_para(doc, "经济回报体现在服务转化、机构合作和复制扩展形成的收入能力。")
    add_subtitle(doc, "（二）社会回报")
    add_para(doc, "社会回报体现在心理支持可及性提升、情绪教育参与度提升和关系沟通质量改善。")
    add_subtitle(doc, "（三）教育与文化回报")
    add_para(doc, "教育回报体现在学校心理健康工作工具升级，文化回报体现在国画审美与疗愈实践的数字化传播。")
    add_subtitle(doc, "（四）评审价值表达")
    add_para(doc, "综合回报分析应强调“可验证、可复制、可持续”三项核心价值，以契合国奖评审对项目长效价值的关注。")
    add_caption(doc, "表13-1 综合回报评估框架（占位，待插入）")

    doc.add_heading("第十四章：发展规划与战略部署", level=1)
    add_subtitle(doc, "（一）近期规划（0-12个月）")
    add_para(doc, "重点完成样本扩展、核心模块稳定、校内试点闭环，形成首轮标准化流程。")
    add_subtitle(doc, "（二）中期规划（1-3年）")
    add_para(doc, "重点推进跨校复制和机构共建，形成多场景适配能力与合作网络。")
    add_subtitle(doc, "（三）远期规划（3年以上）")
    add_para(doc, "重点打造平台化能力，推动标准输出和生态合作，构建可持续增长路径。")
    add_caption(doc, "图14-1 三阶段战略部署路线图（占位，待插入）")

    doc.add_heading("第十五章：结论与合作诉求", level=1)
    add_subtitle(doc, "（一）项目结论")
    add_para(doc, "本项目已形成问题导向明确、技术路径清晰、证据链条完整的参赛方案，具备国奖评审所关注的创新性、实证性和可复制性。")
    add_subtitle(doc, "（二）合作诉求")
    add_para(doc, "下一阶段希望在试点场景、专家评估、跨学科共建和数据协同方面获得支持，进一步提升项目验证强度与落地质量。")

    doc.add_heading("附录A：市场调研详细数据", level=1)
    add_para(doc, "建议附上问卷原题、样本结构、统计口径、原始数据表和关键图表，并在图注中明确样本范围、时间窗口与计算方式。")
    add_caption(doc, "表A-1 问卷题项与指标映射表（占位，待插入）")
    add_caption(doc, "图A-1 样本分布图（占位，待插入）")

    doc.add_heading("附录B：技术方案详细文档", level=1)
    add_para(doc, "建议附上架构图、模块说明、数据流图、接口说明、异常处理机制和版本日志，确保技术实现路径可审阅、可复核。")
    add_caption(doc, "图B-1 详细技术架构图（占位，待插入）")
    add_caption(doc, "表B-1 模块接口清单（占位，待插入）")

    out = r"c:\Users\33788\Desktop\moxin\三创赛项目计划书_国奖冲刺版_详细扩写_墨韵心调.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
