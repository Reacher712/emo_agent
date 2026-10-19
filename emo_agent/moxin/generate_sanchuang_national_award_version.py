from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def set_layout(doc: Document):
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(2.5)
        sec.header_distance = Cm(1.5)
        sec.footer_distance = Cm(1.5)

        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.text = "三创赛国奖冲刺版计划书（墨韵心调）"
        if hp.runs:
            run = hp.runs[0]
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)

        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.text = "第 "
        add_page_number(fp)
        fp.add_run(" 页")
        for run in fp.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)


def set_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(0.74)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "黑体"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h1.font.size = Pt(16)
    h1.font.bold = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "黑体"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h2.font.size = Pt(14)
    h2.font.bold = True


def add_para(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if p.runs:
        run = p.runs[0]
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("三创赛项目计划书（国奖冲刺版）")
    r.bold = True
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    r.font.size = Pt(24)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("墨韵心调——AI驱动的国画多模态情感疗愈平台")
    r2.font.name = "宋体"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r2.font.size = Pt(15)

    doc.add_paragraph("")
    for line in [
        "版本说明：创新-验证-可复制导向结构",
        "适用场景：三创赛国奖评审答辩",
        "日期：2026年3月",
    ]:
        pp = doc.add_paragraph(line)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_toc(doc: Document):
    doc.add_heading("目录", level=1)
    toc = [
        "第一章：执行摘要",
        "第二章：项目背景与问题定义",
        "第三章：市场需求与用户证据",
        "第四章：竞争格局与机会窗口",
        "第五章：产品与服务方案",
        "第六章：技术架构与实现路径",
        "第七章：创新点与知识产权布局（前置）",
        "第八章：实证验证与阶段成果（重点）",
        "第九章：运营模式与推广策略",
        "第十章：组织架构与实施计划",
        "第十一章：风险、伦理与合规治理（前置）",
        "第十二章：财务预测与分析",
        "第十三章：综合回报分析（经济+社会+教育）",
        "第十四章：发展规划与战略部署",
        "第十五章：结论与合作诉求",
        "附录A：市场调研详细数据",
        "附录B：技术方案详细文档",
    ]
    for line in toc:
        doc.add_paragraph(line)
    doc.add_page_break()


def main():
    doc = Document()
    set_layout(doc)
    set_styles(doc)
    add_cover(doc)
    add_toc(doc)

    doc.add_heading("第一章：执行摘要", level=1)
    add_para(doc, "本项目聚焦青年群体情绪压力管理与关系沟通支持，提出以国画为文化核心、以多智能体为技术底座、以动态追踪为评估机制的多模态情感疗愈方案。项目目标是构建可进入、可持续、可验证的数字心理支持路径，并形成可复制的校园示范模式。")
    add_para(doc, "阶段性成果显示，项目已完成方案构建、数据脚本、可视化看板和路演材料，并形成需求调研（n=180）与试用评估（n=60）数据基础，具备进一步扩大样本验证和场景试点的条件。")

    doc.add_heading("第二章：项目背景与问题定义", level=1)
    add_para(doc, "在政策与行业双重驱动下，心理健康服务数字化与传统文化创造性转化呈现融合趋势。当前青年群体对于低门槛、强隐私、可持续的情绪支持工具需求显著增长。")
    add_para(doc, "现有方案存在三类不足：一是专业服务触达频率有限；二是通用情绪工具文化适配度不足；三是AI内容工具缺乏干预闭环与效果评估。项目据此提出系统性解决路径。")
    add_caption(doc, "图2-1 问题定义与价值目标映射图（占位，待插入）")

    doc.add_heading("第三章：市场需求与用户证据", level=1)
    add_para(doc, "基于调研数据，近1个月压力较高及以上占比为67.2%，显示目标用户群体存在明显的日常调节需求。愿意尝试匿名线上心理支持占比77.2%，愿意参加艺术化情绪调节活动占比81.7%，说明项目路径具有较高用户接受度。")
    add_para(doc, "需求侧证据支持项目优先建设三类能力：个性化干预能力、趋势追踪能力、关系场景沟通能力。")
    add_caption(doc, "表3-1 样本结构与核心指标汇总（占位，待插入）")
    add_caption(doc, "图3-1 需求证据可视化图（占位，待插入）")

    doc.add_heading("第四章：竞争格局与机会窗口", level=1)
    add_para(doc, "竞品格局可分为冥想类应用、心理咨询平台与AI生成工具。前两者分别在轻量化与专业化上具有优势，但在“文化疗愈+动态追踪+关系场景”方面尚未形成完整闭环；后者具备生成能力，但缺乏干预机制与评估体系。")
    add_para(doc, "本项目通过组合式能力构建实现差异化定位，在产品体验、验证机制与场景延展性方面形成竞争优势。")
    add_caption(doc, "图4-1 竞品定位矩阵（占位，待插入）")

    doc.add_heading("第五章：产品与服务方案", level=1)
    add_para(doc, "产品方案由“个人疗愈服务”与“关系疗愈服务”构成。个人疗愈服务覆盖情绪识别、国画生成、音乐匹配和情绪看板；关系疗愈服务通过双人共绘与关系温度报告支持沟通优化。")
    add_para(doc, "服务设计强调“低认知负担、可持续复访、结果可解释”，通过过程记录与反馈机制提升用户长期使用价值。")
    add_caption(doc, "图5-1 产品服务全景图（占位，待插入）")

    doc.add_heading("第六章：技术架构与实现路径", level=1)
    add_para(doc, "系统采用“用户层—协调层—智能体层—数据层”四层架构。协调层负责任务编排、状态管理与反馈优化，保障多模块协同效率与一致性。")
    add_para(doc, "实现路径采取分阶段推进：先完成数据与展示链路，再推进核心模块联调，最后进入试点验证与迭代优化。")
    add_caption(doc, "图6-1 系统架构图（占位，待插入）")
    add_caption(doc, "图6-2 关键流程时序图（占位，待插入）")

    doc.add_heading("第七章：创新点与知识产权布局（前置）", level=1)
    add_para(doc, "创新点一：多智能体协同机制，提升任务拆解与内容联动能力；创新点二：国画情感量化机制，提升干预内容可解释性；创新点三：动态情绪追踪机制，形成长期反馈闭环；创新点四：双人共绘关系分析机制，拓展关系疗愈应用边界。")
    add_para(doc, "知识产权布局建议围绕软件著作权、流程方法文档、关键模块专有实现与品牌标识开展，形成技术与品牌双重保护。")
    add_caption(doc, "表7-1 创新点-壁垒-知识产权对应表（占位，待插入）")

    doc.add_heading("第八章：实证验证与阶段成果（重点）", level=1)
    add_para(doc, "本章为评审重点章节。项目采用前后测与体验评估结合的方法，围绕DASS-21、主观平静感、情绪可控感及可用性指标开展验证。")
    add_para(doc, "阶段性结果显示，压力、焦虑、抑郁指标呈下降趋势，平静感与可控感提升，SUS处于良好区间。上述结果初步验证了方案可行性与应用潜力。")
    add_para(doc, "同时，本项目形成了可复核的阶段成果清单，包括问卷数据脚本、可视化看板、竞品证据页与路演材料，为后续规模化试点提供支撑。")
    add_caption(doc, "表8-1 验证指标体系（占位，待插入）")
    add_caption(doc, "图8-1 DASS前后测结果图（占位，待插入）")
    add_caption(doc, "图8-2 可用性评估结果图（占位，待插入）")

    doc.add_heading("第九章：运营模式与推广策略", level=1)
    add_para(doc, "运营策略采用“校内试点—跨校复制—机构协作”分层推进模式。前期聚焦场景验证与用户口碑，中期强化标准化实施包建设，后期推动多场景协同扩展。")
    add_para(doc, "推广策略强调“证据传播+场景体验+合作背书”三位一体，提高项目可信度与转化效率。")
    add_caption(doc, "图9-1 运营与推广路径图（占位，待插入）")

    doc.add_heading("第十章：组织架构与实施计划", level=1)
    add_para(doc, "团队由技术、产品、心理、艺术与运营方向组成，具备跨学科协同能力。实施计划按照阶段里程碑推进，设置可量化交付项与过程复盘机制。")
    add_para(doc, "项目管理采用周度执行跟踪与月度目标评估结合模式，保障计划执行质量与风险前置识别能力。")
    add_caption(doc, "图10-1 团队组织结构图（占位，待插入）")
    add_caption(doc, "表10-1 里程碑与交付项清单（占位，待插入）")

    doc.add_heading("第十一章：风险、伦理与合规治理（前置）", level=1)
    add_para(doc, "针对心理+AI场景特征，本项目建立了“隐私保护—内容安全—边界管理—转介机制”四位一体治理框架。")
    add_para(doc, "隐私层面采用最小化采集、匿名化处理与访问审计；内容层面引入风险识别与人工复核；边界层面明确项目非医疗诊断定位；高风险场景下提供人工转介建议。")
    add_caption(doc, "表11-1 风险与合规治理矩阵（占位，待插入）")

    doc.add_heading("第十二章：财务预测与分析", level=1)
    add_para(doc, "财务预测建议采用保守、中性、积极三情景模型，重点评估研发投入、运维成本、推广成本与收入结构变化。")
    add_para(doc, "通过分阶段预算管理和滚动复盘机制，确保资源配置与阶段目标保持一致。")
    add_caption(doc, "表12-1 财务预测参数与情景假设（占位，待插入）")

    doc.add_heading("第十三章：综合回报分析（经济+社会+教育）", level=1)
    add_para(doc, "国奖评审强调综合价值。本项目回报分析包含经济回报、社会效益、教育效益与文化传承效益四维度。")
    add_para(doc, "经济回报体现在服务转化与合作拓展；社会效益体现在心理健康支持可及性提升；教育效益体现在校园心理服务工具升级；文化效益体现在国画疗愈路径的数字化传播。")
    add_caption(doc, "表13-1 综合回报评估框架（占位，待插入）")

    doc.add_heading("第十四章：发展规划与战略部署", level=1)
    add_para(doc, "发展规划分为近期、中期、远期三阶段：近期完成样本扩展与产品稳定，中期形成跨校复制能力，远期推进平台化与生态协同。")
    add_para(doc, "战略部署遵循“创新机制沉淀—标准流程输出—合作网络拓展”的路径，提升项目长期竞争力。")
    add_caption(doc, "图14-1 三阶段战略部署图（占位，待插入）")

    doc.add_heading("第十五章：结论与合作诉求", level=1)
    add_para(doc, "综上，本项目已形成明确问题导向、可验证技术路径与可扩展应用场景，具备国奖评审所关注的创新性、实证性与可复制性。")
    add_para(doc, "后续希望在试点学校资源、专家评估支持、跨学科协同机制等方面获得进一步合作，共同推动项目高质量落地与持续迭代。")

    doc.add_heading("附录A：市场调研详细数据", level=1)
    add_para(doc, "建议附上问卷题项、样本结构、统计口径说明、关键图表及原始表格，确保结论可复核、可追溯。")

    doc.add_heading("附录B：技术方案详细文档", level=1)
    add_para(doc, "建议附上系统架构图、模块流程图、数据流设计、接口说明、关键算法说明与版本记录，提升评审对技术可行性的判断效率。")

    out = r"c:\Users\33788\Desktop\moxin\三创赛项目计划书_国奖冲刺版_墨韵心调.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
