from docx import Document
import glob
import os


def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def main():
    files = glob.glob(r"c:/Users/33788/Desktop/moxin/*.docx")
    if not files:
        raise FileNotFoundError("未找到 docx 文件。")

    target = max(files, key=os.path.getmtime)
    doc = Document(target)

    # 锚点：正文第一章（后面紧跟【1.1 ...】）
    idx_body_ch1 = None
    for i in range(len(doc.paragraphs) - 1):
        cur = (doc.paragraphs[i].text or "").strip()
        nxt = (doc.paragraphs[i + 1].text or "").strip()
        if cur == "第一章 执行摘要" and nxt.startswith("【1.1"):
            idx_body_ch1 = i
            break

    if idx_body_ch1 is None:
        raise RuntimeError("未定位到正文第一章锚点。")

    chapter_para = doc.paragraphs[idx_body_ch1]

    # 清理文档前导区（保留前4段：标题、副标题、空行、使用说明）
    # 然后完整重建“使用说明补充 + 目录（含二级小点）”
    start_clean = 4
    for i in range(idx_body_ch1 - 1, start_clean - 1, -1):
        delete_paragraph(doc.paragraphs[i])

    rebuilt_front = [
        "1) 本模板已包含“目录 + 每章填空提示 + 示例文案”。",
        "2) 方括号内容（如[请填写]）表示你需要替换。",
        "3) 示例文案可直接改写后使用，建议补充你们最新数据与截图。",
        "",
        "目录（建议定稿后再微调页码）",
        "第一章 执行摘要",
        "  1.1 项目一句话定位",
        "  1.2 核心价值主张",
        "  1.3 当前进展与阶段成果",
        "第二章 项目背景与问题定义",
        "  2.1 政策与行业背景",
        "  2.2 目标问题与痛点拆解",
        "  2.3 目标用户与核心场景",
        "第三章 市场与用户调研",
        "  3.1 调研方法与样本说明",
        "  3.2 核心调研结论",
        "  3.3 需求优先级与机会窗口",
        "  3.4 竞品分析与差异化定位",
        "第四章 解决方案与产品设计",
        "  4.1 产品总体方案与业务闭环",
        "  4.2 功能模块设计",
        "  4.3 用户旅程与关键交互流程",
        "第五章 技术方案与实现路径",
        "  5.1 系统架构设计",
        "  5.2 关键机制与算法逻辑",
        "  5.3 开发进度与部署计划",
        "第六章 创新点与核心竞争力",
        "  6.1 创新点一：智能体协同",
        "  6.2 创新点二：国画情感量化",
        "  6.3 创新点三：动态情绪追踪",
        "  6.4 创新点四：双人共绘与关系温度",
        "第七章 商业模式与运营推广",
        "  7.1 商业模式设计",
        "  7.2 推广路径与合作策略",
        "  7.3 核心运营指标",
        "第八章 实验验证与阶段成果",
        "  8.1 验证方法与评价指标",
        "  8.2 前后测结果与效果分析",
        "  8.3 阶段成果清单",
        "  8.4 局限性与优化计划",
        "第九章 风险控制、伦理与合规",
        "  9.1 数据隐私与安全治理",
        "  9.2 内容安全与应用边界",
        "  9.3 风险识别与应对预案",
        "第十章 团队组织与实施计划",
        "  10.1 团队结构与职责分工",
        "  10.2 里程碑计划与时间安排",
        "  10.3 资源保障与外部协同",
        "第十一章 财务预算与资源需求",
        "  11.1 预算结构与成本构成",
        "  11.2 资金使用原则",
        "  11.3 预期产出与投入产出关系",
        "第十二章 结论与合作诉求",
        "  12.1 项目结论",
        "  12.2 下一阶段目标",
        "  12.3 合作诉求",
        "附录 建议附加材料清单",
        "",
    ]

    for line in rebuilt_front:
        p = chapter_para.insert_paragraph_before(line)
        p.style = doc.styles["Normal"]

    doc.save(target)
    print(target)


if __name__ == "__main__":
    main()
