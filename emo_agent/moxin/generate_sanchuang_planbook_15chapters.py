from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def set_layout(doc: Document):
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(2.5)
        sec.header_distance = Cm(1.5)
        sec.footer_distance = Cm(1.5)

        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.text = "三创赛项目计划书（墨韵心调）"
        hrun = hp.runs[0]
        hrun.font.name = "宋体"
        hrun._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        hrun.font.size = Pt(10.5)

        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.text = "第 "
        add_page_number(fp)
        fp.add_run(" 页")
        for run in fp.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)


def set_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(0.74)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "黑体"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h1.font.size = Pt(16)
    h1.font.bold = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "黑体"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h2.font.size = Pt(14)
    h2.font.bold = True


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("三创赛项目计划书")
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    r.font.size = Pt(24)
    r.bold = True

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("墨韵心调——AI驱动的国画多模态情感疗愈平台")
    r2.font.name = "宋体"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r2.font.size = Pt(15)

    doc.add_paragraph("")
    for line in [
        "参赛类型：三创赛创新创业项目",
        "文档版本：结构化计划书（15章）",
        "日期：2026年3月",
    ]:
        pp = doc.add_paragraph(line)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_toc(doc: Document):
    doc.add_heading("目录", level=1)
    toc = [
        "第一部分：项目概述与市场分析",
        "  第一章：封面与摘要",
        "  第二章：项目背景与意义",
        "  第三章：市场分析",
        "  第四章：竞争分析",
        "第二部分：产品服务与技术方案",
        "  第五章：产品与服务介绍",
        "  第六章：技术架构与实现方案",
        "  第七章：创新点与知识产权",
        "  第八章：运营模式与盈利模式",
        "第三部分：运营管理与团队建设",
        "  第九章：组织架构与团队介绍",
        "  第十章：运营管理与实施计划",
        "  第十一章：市场营销与推广策略",
        "第四部分：财务分析与发展规划",
        "  第十二章：财务预测与分析",
        "  第十三章：投资回报分析",
        "  第十四章：发展规划与战略部署",
        "  第十五章：风险分析与应对措施",
        "附录",
        "  附录A：市场调研详细数据",
        "  附录B：技术方案详细文档",
    ]
    for line in toc:
        doc.add_paragraph(line)
    doc.add_page_break()


def add_para(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if p.runs:
        run = p.runs[0]
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def main():
    doc = Document()
    set_layout(doc)
    set_styles(doc)
    add_cover(doc)
    add_toc(doc)

    doc.add_heading("第一部分：项目概述与市场分析", level=1)

    doc.add_heading("第一章：封面与摘要", level=2)
    add_para(doc, "本项目围绕青年群体高频情绪压力与亲密关系沟通需求，提出“情绪识别—国画生成—音乐匹配—反馈追踪”一体化多模态情感疗愈方案。项目以中国传统国画为核心媒介，在数字化交互场景中实现心理支持的日常化、低门槛化与可持续化。")
    add_para(doc, "现阶段项目已形成“方案文档+问卷数据+可视化看板+路演材料”的完整参赛资料链，并完成阶段性样本验证，为后续校内试点和跨场景推广提供基础。")

    doc.add_heading("第二章：项目背景与意义", level=2)
    add_para(doc, "在政策层面，心理健康教育与美育融合持续受到重视；在技术层面，生成式人工智能显著提升了个性化内容生产效率；在场景层面，青年群体对隐私友好、使用便捷、文化认同感强的情绪支持工具存在现实需求。")
    add_para(doc, "项目意义体现在三个方面：其一，补足专业心理服务之外的日常支持环节；其二，促进传统文化在数字心理健康场景中的创新转化；其三，为高校心理服务提供可视化、可追踪、可评估的辅助工具。")
    add_caption(doc, "图2-1 项目问题-方案-价值逻辑图（占位，待插入）")

    doc.add_heading("第三章：市场分析", level=2)
    add_para(doc, "基于现有调研数据（需求样本n=180），目标用户主要为15-22岁青年群体，其中“近1个月压力较高及以上”占比达到67.2%，显示出高频轻干预的显著需求。")
    add_para(doc, "用户对匿名线上支持与艺术化情绪调节接受度较高，相关指标分别达到77.2%与81.7%。该结果表明，兼顾隐私性、可参与性和文化体验的产品路径具有较好的市场适配性。")
    add_para(doc, "从应用落地视角看，项目可先聚焦高校场景形成示范，再向心理机构、文化教育机构等场景延展，具备由点到面的推广潜力。")
    add_caption(doc, "表3-1 用户样本结构与核心指标统计表（占位，待插入）")
    add_caption(doc, "图3-1 目标市场分层图（占位，待插入）")

    doc.add_heading("第四章：竞争分析", level=2)
    add_para(doc, "当前相关产品可分为冥想类应用、心理咨询平台与AI生成工具三类。冥想类应用具备轻量优势，但长期留存和文化差异化相对有限；心理咨询平台专业性强，但存在成本与触达门槛；AI生成工具具备内容能力，但干预闭环与效果评估不足。")
    add_para(doc, "本项目的竞争策略为“文化疗愈内容+情绪趋势追踪+关系共绘分析”组合创新，通过产品机制与数据反馈共同构建差异化壁垒。")
    add_caption(doc, "图4-1 竞品对比矩阵（占位，待插入）")

    doc.add_heading("第二部分：产品服务与技术方案", level=1)

    doc.add_heading("第五章：产品与服务介绍", level=2)
    add_para(doc, "产品服务由个人疗愈服务与关系疗愈服务构成。个人疗愈服务覆盖情绪识别、国画生成、音乐匹配和情绪周报；关系疗愈服务通过双人共绘与互动行为分析生成关系温度报告。")
    add_para(doc, "项目坚持“轻交互、重反馈、可持续”原则，强调从一次性体验转向长期陪伴，增强用户实际获得感。")
    add_caption(doc, "图5-1 核心功能模块图（占位，待插入）")

    doc.add_heading("第六章：技术架构与实现方案", level=2)
    add_para(doc, "系统采用四层架构：用户交互层、智能体协调层、智能体能力层、数据支撑层。协调层负责任务拆解、状态管理与反馈回流，保障多模块协同效率与一致性。")
    add_para(doc, "在实现路径上，项目先完成数据与展示链路，再推进模块联调和试点部署，形成“可演示—可验证—可推广”的分阶段技术路线。")
    add_caption(doc, "图6-1 系统技术架构图（占位，待插入）")
    add_caption(doc, "图6-2 关键流程时序图（占位，待插入）")

    doc.add_heading("第七章：创新点与知识产权", level=2)
    add_para(doc, "创新点主要包括：智能体协同机制、国画情感量化机制、动态情绪轨迹追踪机制、双人共绘关系温度分析机制。上述创新点在内容生成、过程解释与结果评估三个层面形成协同。")
    add_para(doc, "知识产权方面，建议围绕软件著作权、算法流程文档、数据处理方法说明等方向开展布局，并同步完善版本留痕和成果归档。")
    add_caption(doc, "表7-1 创新点与知识产权布局表（占位，待插入）")

    doc.add_heading("第八章：运营模式与盈利模式", level=2)
    add_para(doc, "运营模式采用“校内试点—跨校复制—机构合作”三级路径，以场景验证驱动产品迭代。前期重在公益化和示范化应用，中后期通过机构服务和数据服务形成可持续收入结构。")
    add_para(doc, "盈利模式建议采用分层设计：基础服务提升用户规模与口碑，进阶服务面向学校/机构提供管理看板与团体分析能力，形成稳定付费基础。")
    add_caption(doc, "图8-1 运营与盈利模式图（占位，待插入）")

    doc.add_heading("第三部分：运营管理与团队建设", level=1)

    doc.add_heading("第九章：组织架构与团队介绍", level=2)
    add_para(doc, "团队组织遵循跨学科协同原则，覆盖技术研发、产品设计、心理支持、艺术内容与运营推广。该组织方式有助于提升项目的系统完整性与落地效率。")
    add_para(doc, "建议在计划书中明确团队成员角色分工、负责模块、阶段目标与评价指标，增强项目执行可信度。")
    add_caption(doc, "图9-1 团队组织架构图（占位，待插入）")

    doc.add_heading("第十章：运营管理与实施计划", level=2)
    add_para(doc, "实施计划建议按照“准备期—联调期—试点期—优化期”推进，并在每阶段设置可量化里程碑，如功能完成率、样本覆盖数、复访率提升幅度等。")
    add_para(doc, "项目管理采用周度复盘与月度评估结合机制，确保目标可监控、问题可定位、资源可调整。")
    add_caption(doc, "图10-1 实施计划甘特图（占位，待插入）")
    add_caption(doc, "表10-1 阶段里程碑与交付项（占位，待插入）")

    doc.add_heading("第十一章：市场营销与推广策略", level=2)
    add_para(doc, "推广策略建议采用“校园活动触达+线上内容传播+合作机构背书”联动模式，构建品牌认知与用户转化双通道。")
    add_para(doc, "营销内容应突出“文化疗愈、可解释评估、真实数据支持”三项核心信息，避免仅以概念叙事替代证据表达。")
    add_caption(doc, "图11-1 市场推广路径图（占位，待插入）")

    doc.add_heading("第四部分：财务分析与发展规划", level=1)

    doc.add_heading("第十二章：财务预测与分析", level=2)
    add_para(doc, "财务预测应分短期（1年）与中期（3年）进行，重点关注研发投入、运营成本、获客成本与服务收入变化趋势。")
    add_para(doc, "建议采用保守、中性、积极三种情景进行敏感性分析，以增强财务预测的稳健性和可解释性。")
    add_caption(doc, "表12-1 财务预测模型与参数说明（占位，待插入）")

    doc.add_heading("第十三章：投资回报分析", level=2)
    add_para(doc, "投资回报分析建议覆盖资金使用效率、阶段产出价值、用户增长贡献及合作转化效果。对于大创和竞赛场景，可同步展示社会效益与教育效益，形成综合回报论证。")
    add_para(doc, "在指标选择上，建议加入单位投入产出、阶段复用率、试点复制成本等指标，体现项目可持续发展能力。")
    add_caption(doc, "表13-1 投资回报测算表（占位，待插入）")

    doc.add_heading("第十四章：发展规划与战略部署", level=2)
    add_para(doc, "战略部署建议分为近期（样本扩展与产品稳定）、中期（跨校复制与机构共建）、远期（平台化能力与生态合作）三个阶段。")
    add_para(doc, "发展规划应坚持“技术迭代—场景验证—标准化输出”路径，以形成具备可迁移能力的项目方法论。")
    add_caption(doc, "图14-1 三阶段发展路线图（占位，待插入）")

    doc.add_heading("第十五章：风险分析与应对措施", level=2)
    add_para(doc, "主要风险包括技术风险、数据风险、合规风险和运营风险。技术风险表现为模型稳定性和偏差问题；数据风险表现为样本代表性不足；合规风险表现为隐私与边界管理；运营风险表现为早期留存与转化不确定性。")
    add_para(doc, "应对措施包括：建立人工校验与模型迭代机制、完善匿名化与权限审计、明确高风险场景人工转介流程、通过小规模试点逐步放大推广范围。")
    add_caption(doc, "表15-1 风险识别与应对矩阵（占位，待插入）")

    doc.add_heading("附录", level=1)
    doc.add_heading("附录A：市场调研详细数据", level=2)
    add_para(doc, "建议附上问卷题项、样本结构、统计口径、原始结果表及关键图表，确保结论具备可追溯性。")
    add_para(doc, "推荐纳入文件：青少年问卷结果汇总、压力分布分层、DASS前后测、可用性评估结果。")

    doc.add_heading("附录B：技术方案详细文档", level=2)
    add_para(doc, "建议附上系统架构图、模块说明、数据流说明、接口清单、关键算法说明及版本记录，便于评审理解技术可行性与工程完整性。")
    add_para(doc, "推荐纳入文件：总览文档、技术流程图、看板页面说明、竞品证据引用清单。")

    out = r"c:\Users\33788\Desktop\moxin\三创赛项目计划书_15章结构版_墨韵心调.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
